#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Actualiza la columna VOD del Excel con las plataformas donde cada película
está disponible EN ESPAÑA, según TMDB (datos de JustWatch).

USO:
    export TMDB_API_KEY="tu_clave"
    python3 actualizar_vod.py

REQUISITOS:
    pip3 install requests openpyxl

ARCHIVOS QUE NECESITA (en la misma carpeta):
    - Movie Data Base by Keizar.xlsx
    - tmdb_ids.csv
"""

import csv
import os
import re
import sys
import time
import unicodedata
from datetime import date

import requests
import openpyxl
from openpyxl.styles import Alignment, Font

# ---------------------------------------------------------------- configuración
EXCEL = "Movie Data Base by Keizar.xlsx"
HOJA = "Movie Data Base by Keizar"
IDS_CSV = "tmdb_ids.csv"
COL_TITULO = 1
COL_ANIO = 2
COL_VOD = 12
REGION = "ES"
PAUSA = 0.06          # segundos entre llamadas (TMDB admite ~50/s)
TIPOS = ("flatrate", "free", "ads")   # solo suscripción/gratis, no alquiler

API_KEY = os.environ.get("TMDB_API_KEY", "").strip()
if not API_KEY:
    sys.exit("ERROR: define la variable TMDB_API_KEY antes de ejecutar.")

# Nombres cortos para que quepan en la celda
ABREVIA = {
    "Netflix": "Netflix",
    "Amazon Prime Video": "Prime",
    "Amazon Prime Video with Ads": "Prime",
    "Disney Plus": "Disney+",
    "Movistar Plus": "Movistar",
    "Movistar Plus+": "Movistar",
    "Apple TV+": "Apple TV+",
    "Apple TV Plus": "Apple TV+",
    "Filmin": "Filmin",
    "HBO Max": "HBO Max",
    "Max": "HBO Max",
    "SkyShowtime": "SkyShowtime",
    "Rakuten TV": "Rakuten",
    "Atres Player": "atresplayer",
    "RTVE Play": "RTVE Play",
    "Pluto TV": "Pluto TV",
    "Crunchyroll": "Crunchyroll",
    "FlixOlé": "FlixOlé",
    "Acontra Plus": "Acontra",
    "Tivify": "Tivify",
    "Runtime": "Runtime",
    "Prime Video": "Prime",
}


def norm(s):
    if not s:
        return ""
    s = str(s).lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s)


def anio(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def cargar_ids():
    """Devuelve {(titulo_normalizado, anio): tmdb_id}"""
    mapa = {}
    with open(IDS_CSV, encoding="utf-8-sig") as f:
        for fila in csv.DictReader(f):
            tid = (fila.get("tmdb_id") or "").strip()
            if not tid:
                continue
            mapa[(norm(fila["titulo"]), anio(fila["anio"]))] = tid
    return mapa


def plataformas(sesion, tmdb_id):
    """Consulta TMDB y devuelve la cadena de plataformas, o None si falla."""
    url = f"https://api.themoviedb.org/3/movie/{tmdb_id}/watch/providers"
    try:
        r = sesion.get(url, params={"api_key": API_KEY}, timeout=15)
    except requests.RequestException:
        return None
    if r.status_code == 429:                 # límite de peticiones
        time.sleep(2)
        return plataformas(sesion, tmdb_id)
    if r.status_code != 200:
        return None
    datos = r.json().get("results", {}).get(REGION)
    if not datos:
        return ""                            # no disponible en España
    nombres = []
    for tipo in TIPOS:
        for p in datos.get(tipo, []):
            n = p.get("provider_name", "")
            n = ABREVIA.get(n, n)
            if n and n not in nombres:
                nombres.append(n)
    return ", ".join(nombres)


def main():
    ids = cargar_ids()
    print(f"Identificadores cargados: {len(ids)}")

    wb = openpyxl.load_workbook(EXCEL)
    ms = wb[HOJA]

    sesion = requests.Session()
    con_datos = sin_datos = sin_id = 0

    for i in range(2, ms.max_row + 1):
        titulo = ms.cell(i, COL_TITULO).value
        if not titulo:
            continue
        clave = (norm(titulo), anio(ms.cell(i, COL_ANIO).value))
        tid = ids.get(clave)
        if not tid:
            sin_id += 1
            continue

        valor = plataformas(sesion, tid)
        time.sleep(PAUSA)
        if valor is None:                    # error de red: no tocar la celda
            continue

        c = ms.cell(i, COL_VOD)
        c.value = valor if valor else None
        c.font = Font(name="Calibri", size=11)
        c.alignment = Alignment(wrap_text=False, vertical="center")

        if valor:
            con_datos += 1
        else:
            sin_datos += 1

        if (con_datos + sin_datos) % 200 == 0:
            print(f"  procesadas {con_datos + sin_datos}...")

    # Marca de la última actualización, en la cabecera de la columna
    ms.cell(1, COL_VOD).value = f"VOD ({date.today():%d/%m/%Y})"

    wb.save(EXCEL)
    print(f"\nDisponibles en alguna plataforma: {con_datos}")
    print(f"No disponibles en España: {sin_datos}")
    print(f"Sin identificador TMDB: {sin_id}")
    print(f"Guardado: {EXCEL}")


if __name__ == "__main__":
    main()
