#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Exporta el Excel a datos.json (lo que lee la app del móvil)."""
import json
from datetime import date

import openpyxl

from comun import EXCEL, HOJA_PELIS, HOJA_SERIES, P, S


def limpio(v):
    return None if v in (None, "", 0) else v


def main():
    import csv, os
    fechas = {}
    if os.path.exists("historial_votos.csv"):
        with open("historial_votos.csv", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                fechas[r["imdb_id"]] = r["fecha"]
    posters = {}
    if os.path.exists("posters.csv"):
        with open("posters.csv", encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                posters[r["imdb_id"]] = (r.get("poster") or None, r.get("fondo") or None)
    po = lambda tt: posters.get(tt, (None, None))
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    pelis = []
    for r in wb[HOJA_PELIS].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        g = lambda k: r[P[k] - 1] if P[k] - 1 < len(r) else None
        d = dict(t=g("titulo"), a=g("anio"), d=g("dur"), dir=g("director"), i=g("imdb"), m=g("mi"),
                 v=g("votos"), p=g("pais"), g=g("genero"), sg=g("subgenero"), vod=g("vod"),
                 s=g("sinopsis"), o=g("oscars"), op=g("oscar_peli"), od=g("oscar_dir"), on=g("nom"),
                 f=g("fest"), gl=g("globo"), ss=g("ss"), pt=round(g("puntos") or 0, 1),
                 h=g("hist"), jw=g("jw"), t150=g("top150"), tt=g("tt"), fv=fechas.get(g("tt")), po=po(g("tt"))[0], bd=po(g("tt"))[1])
        pelis.append({k: v for k, v in d.items() if limpio(v) is not None})
    series = []
    for r in wb[HOJA_SERIES].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        g = lambda k: r[S[k] - 1] if S[k] - 1 < len(r) else None
        d = dict(t=g("titulo"), a=g("anio"), af=g("fin"), tipo=g("tipo"), dir=g("creador"), i=g("imdb"),
                 m=g("mi"), v=g("votos"), p=g("pais"), g=g("genero"), sg=g("subgenero"), vod=g("vod"),
                 s=g("sinopsis"), e=g("emmys"), es=g("emmy_serie"), en=g("nom_emmy"), h=g("hist"), tt=g("tt"), fv=fechas.get(g("tt")), po=po(g("tt"))[0], bd=po(g("tt"))[1])
        series.append({k: v for k, v in d.items() if limpio(v) is not None})
    vod_fecha = str(wb[HOJA_PELIS].cell(1, P["vod"]).value or "")
    salida = dict(generado=f"{date.today():%Y-%m-%d}", vod=vod_fecha, pelis=pelis, series=series)
    with open("datos.json", "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, separators=(",", ":"))
    print(f"datos.json: {len(pelis)} películas, {len(series)} series")


if __name__ == "__main__":
    main()
